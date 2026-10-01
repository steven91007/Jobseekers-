"""Markdown digest, Excel export and console table for one run."""

import json
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from rich import box
from rich.console import Console
from rich.table import Table

from . import freshness, language
from .config import REGION_LABELS, REGIONS, REMOTE_EU
from .normalize import is_ai_engineer_title

TIER_RANK = {"ai_native": 0, "ai_heavy": 1, "other": 2}
TIER_LABEL = {"ai_native": "AI-native", "ai_heavy": "AI-heavy", "other": ""}
LANG_FLAG = {"german": "DE req", "dutch": "NL req", "german_nice_to_have": "DE nice",
             "dutch_nice_to_have": "NL nice"}
REGION_ORDER = [*REGIONS, REMOTE_EU]
TOP_DETAILS = 15
# Added to the rank of AI Engineer titles, the role the candidate is focusing on.
FOCUS_BONUS = 5


def _assessment(row) -> dict:
    try:
        return json.loads(row["assessment"]) if row["assessment"] else {}
    except (TypeError, ValueError):
        return {}


def rank_score(fit: int | None, fresh: int | None, focus: bool, weight: float) -> int | None:
    """Fit blended with freshness, plus a bonus for AI Engineer titles. None when unscored."""
    if fit is None:
        return None
    blended = (1 - weight) * fit + weight * (fresh or 0) + (FOCUS_BONUS if focus else 0)
    return min(100, round(blended))


def non_english_reason(title: str, description: str, extra: dict, assessment: dict) -> str:
    """Why a stored job is not an English-only role, or '' if it is."""
    if extra.get("language_skip"):
        return extra["language_skip"]
    if assessment.get("local_language_required") in ("german", "dutch"):
        return f"requires {assessment['local_language_required'].title()} (scorer)"
    c = language.check(title, description)
    return "" if c.english else c.reason


def build_rows(db_rows, *, window_hours: float, max_age_hours: float, half_life_hours: float = 24.0,
               freshness_weight: float = 0.3, now=None, english_only: bool = False,
               hidden: dict | None = None) -> list[dict]:
    """Report rows for jobs posted within max_age_hours. Older or undated jobs are dropped.

    With english_only, German/Dutch postings and roles requiring those languages are
    dropped too; `hidden` (if given) collects {reason: count} for the report header.
    """
    now = now or freshness.now_utc()
    rows = []
    for r in db_rows:
        posted_at = r["posted_at"] or ""
        if not freshness.within_hours(posted_at, max_age_hours, now):
            continue
        a = _assessment(r)
        extra = json.loads(r["extra"] or "{}")
        if english_only:
            reason = non_english_reason(r["title"], r["description"] or "", extra, a)
            if reason:
                if hidden is not None:
                    hidden[reason] = hidden.get(reason, 0) + 1
                continue
        fresh = freshness.score(posted_at, half_life_hours, now)
        focus = is_ai_engineer_title(r["title"])
        rows.append({
            "job_key": r["job_key"], "title": r["title"], "company": r["company"],
            "tier": r["company_tier"], "location": r["location"], "region": r["region"],
            "url": r["url"], "posted": posted_at[:10], "posted_at": posted_at,
            "age": freshness.label(posted_at, now), "age_hours": freshness.age_hours(posted_at, now),
            "fresh": fresh, "focus": focus, "source": r["source"],
            "is_new": bool(r["is_new"]),
            "recent": freshness.within_hours(posted_at, window_hours, now),
            "score": r["fit_score"], "priority": r["apply_priority"] or "",
            "rank": rank_score(r["fit_score"], fresh, focus, freshness_weight),
            "lang": a.get("local_language_required", ""),
            "visa": a.get("visa_or_relocation", ""),
            "salary": a.get("salary") or extra.get("salary") or "",
            "seniority": a.get("seniority", ""),
            "employer_type": a.get("employer_type", ""),
            "why_fit": a.get("why_fit", ""), "pitch": a.get("suggested_pitch", ""),
            "gaps": a.get("gaps", []), "must_haves": a.get("must_haves", []),
            "trace_id": r["a_trace_id"],
        })
    rows.sort(key=sort_key)
    return rows


def sort_key(row: dict):
    """Scored jobs by rank; then unscored ones, AI Engineer titles first, freshest first."""
    rank = row["rank"]
    fresh = row["fresh"] if row["fresh"] is not None else -1
    return (rank is None, -(rank or 0), not row["focus"], -fresh,
            TIER_RANK.get(row["tier"], 3), _neg_date(row["posted_at"]))


def _neg_date(iso_date: str) -> str:
    # Sort newest first inside an ascending sort by inverting each digit.
    return "".join(chr(ord("9") - ord(c) + ord("0")) if c.isdigit() else c for c in iso_date) or "~"


def _md(text) -> str:
    return str(text or "").replace("|", "\\|").replace("\n", " ").strip()


def _num(value) -> str:
    return "" if value is None else str(value)


def _md_table(rows: list[dict]) -> list[str]:
    out = ["| Rank | Fit | Fresh | New | Posted | Role | Company | Location | Lang | Source |",
           "|---:|---:|---:|:---:|---|---|---|---|---|---|"]
    for r in rows:
        rank = _num(r["rank"])
        if rank and r["priority"] == "now":
            rank = f"**{rank}**"
        company = _md(r["company"]) + (f" · {TIER_LABEL[r['tier']]}" if TIER_LABEL.get(r["tier"]) else "")
        role = f"[{_md(r['title'])}]({r['url']})" + (" 🎯" if r["focus"] else "")
        out.append(
            f"| {rank} | {_num(r['score'])} | {_num(r['fresh'])} | {'🆕' if r['is_new'] else ''} | "
            f"{r['age']} | {role} | {company} | {_md(r['location'])[:40]} | "
            f"{LANG_FLAG.get(r['lang'], '')} | {r['source']} |"
        )
    return out


def _language_note(stats: dict) -> list[str]:
    dropped = {**stats.get("non_english_dropped", {})}
    for reason, n in stats.get("non_english_hidden", {}).items():
        dropped[reason] = dropped.get(reason, 0) + n
    if "non_english_dropped" not in stats and "non_english_hidden" not in stats:
        return []
    total = sum(dropped.values())
    detail = f" ({', '.join(f'{n} {r}' for r, n in sorted(dropped.items(), key=lambda x: -x[1]))})" if total else ""
    return [f"English only: {total} postings in German/Dutch or requiring those languages are not listed{detail}.", ""]


def write_markdown(path: Path, rows: list[dict], *, since: str, max_age_days: int, stats: dict,
                   briefing: str, candidates: list, trace_link: str | None) -> None:
    today = date.today().isoformat()
    scored = [r for r in rows if r["score"] is not None]
    recent = [r for r in rows if r["recent"]]
    lines = [
        f"# AI jobs in Germany, the Netherlands and Dublin: {today}",
        "",
        f"{len(recent)} roles posted in the last {since}, {len(rows)} within {max_age_days} days "
        f"({sum(r['is_new'] for r in rows)} new since the last run, {len(scored)} scored). "
        f"Nothing older than {max_age_days} days is listed.",
        "",
        *_language_note(stats),
        "Rank blends the fit score with freshness and adds a small bonus for AI Engineer titles (🎯). "
        "Fresh is 100 for a posting from right now and halves every day.",
        "",
    ]
    if trace_link:
        lines += [f"Langfuse trace: {trace_link}", ""]

    if briefing:
        lines += ["## Agent briefing", "", briefing.strip(), ""]

    top = [r for r in scored if r["priority"] == "now"][:12]
    if top:
        lines += ["## Apply now", ""] + _md_table(top) + [""]

    for region in REGION_ORDER:
        region_rows = [r for r in rows if r["region"] == region]
        if not region_rows:
            continue
        in_window = [r for r in region_rows if r["recent"]]
        earlier = [r for r in region_rows if not r["recent"]]
        lines += [f"## {REGION_LABELS[region]} ({len(region_rows)})", ""]
        lines += [f"### Posted in the last {since} ({len(in_window)})", ""]
        lines += (_md_table(in_window) if in_window else ["None in this window."]) + [""]
        if earlier:
            lines += [f"### Posted earlier, within {max_age_days} days ({len(earlier)})", ""]
            lines += _md_table(earlier) + [""]

    detailed = [r for r in rows if r["score"] is not None][:TOP_DETAILS]
    if detailed:
        lines += ["## Assessments for the top matches", ""]
        for r in detailed:
            lines += [
                f"### {r['rank']} · {_md(r['title'])} at {_md(r['company'])}",
                "",
                f"Fit {r['score']} · fresh {r['fresh']} · {r['location']} · posted {r['age']} · "
                f"{r['seniority']} · "
                f"{r['employer_type']} · [{r['job_key']}]({r['url']})",
                "",
                f"**Why:** {r['why_fit']}",
                "",
            ]
            if r["must_haves"]:
                lines.append("**Must-haves:** " + "; ".join(r["must_haves"]))
            if r["gaps"]:
                lines.append("**Gaps:** " + "; ".join(r["gaps"]))
            if r["salary"]:
                lines.append(f"**Salary:** {r['salary']}")
            lines += ["", "**Pitch:**", "", *[f"> {line}" for line in r["pitch"].splitlines() if line.strip()], ""]

    if candidates:
        lines += ["## Companies the agent proposes adding to the watchlist", "",
                  "| Company | ATS | Slug | Regions | Why | Careers |", "|---|---|---|---|---|---|"]
        for c in candidates:
            lines.append(f"| {_md(c['name'])} | {c['ats']} | {c['slug']} | {c['regions']} | "
                         f"{_md(c['reason'])} | {c['careers_url']} |")
        lines.append("")

    src = stats.get("sources", {})
    lines += ["## Run details", "",
              f"- Sources: {json.dumps(src)}",
              f"- Failed sources: {', '.join(stats.get('failed_sources', [])) or 'none'}",
              f"- Searches that hit the per-query cap (window may hold more): "
              f"{', '.join(stats.get('capped_searches', [])) or 'none'}",
              f"- Scoring: {json.dumps(stats.get('scoring', {}))}",
              f"- Agent: {json.dumps(stats.get('agent', {}))}", ""]
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text("\n".join(lines), encoding="utf-8")


EXCEL_COLUMNS = [
    ("Region", 12), ("Rank", 7), ("Fit", 6), ("Fresh", 7), ("Priority", 9), ("New", 6),
    ("Posted", 11), ("Posted at", 20), ("AI Engineer", 11), ("Title", 45),
    ("Company", 22), ("Tier", 10), ("Location", 28), ("Local language", 14), ("Visa", 18),
    ("Salary", 20), ("Why fit", 60), ("Source", 10), ("Job key", 22), ("Link", 50),
]


def write_excel(path: Path, rows: list[dict], candidates: list) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Jobs"
    ws.append([c for c, _ in EXCEL_COLUMNS])
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
    for r in rows:
        ws.append([
            REGION_LABELS.get(r["region"], r["region"]), r["rank"], r["score"], r["fresh"],
            r["priority"], "yes" if r["is_new"] else "", r["age"], r["posted_at"],
            "yes" if r["focus"] else "", r["title"], r["company"],
            TIER_LABEL.get(r["tier"], ""), r["location"], r["lang"], r["visa"], r["salary"],
            r["why_fit"], r["source"], r["job_key"], r["url"],
        ])
        link = ws.cell(row=ws.max_row, column=len(EXCEL_COLUMNS))
        if r["url"]:
            link.hyperlink = r["url"]
            link.style = "Hyperlink"
        why_col = [c for c, _ in EXCEL_COLUMNS].index("Why fit") + 1
        ws.cell(row=ws.max_row, column=why_col).alignment = Alignment(wrap_text=True, vertical="top")
    for i, (_, width) in enumerate(EXCEL_COLUMNS, 1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    if candidates:
        cs = wb.create_sheet("Company candidates")
        cs.append(["Company", "ATS", "Slug", "Regions", "Reason", "Careers URL", "Suggested"])
        for cell in cs[1]:
            cell.font = Font(bold=True)
        for c in candidates:
            cs.append([c["name"], c["ats"], c["slug"], c["regions"], c["reason"],
                       c["careers_url"], c["suggested_at"][:10]])
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def print_console(rows: list[dict], console: Console, since: str = "24h") -> None:
    """Every job inside the search window, best rank first; earlier ones are only counted."""
    shown = [r for r in rows if r["recent"]]
    table = Table(title=f"Jobs posted in the last {since} ({len(shown)})", box=box.ROUNDED,
                  header_style="bold cyan", show_lines=False)
    for col, kw in [("#", {"justify": "right", "style": "dim"}), ("Rank", {"justify": "right"}),
                    ("Fit", {"justify": "right"}), ("Fresh", {"justify": "right"}),
                    ("New", {"justify": "center"}), ("Region", {}), ("Posted", {"style": "dim"}),
                    ("Title", {"style": "bold white", "max_width": 44}),
                    ("Company", {"style": "yellow", "max_width": 22}),
                    ("Job key", {"style": "dim", "max_width": 24})]:
        table.add_column(col, **kw)
    for i, r in enumerate(shown, 1):
        rank = _num(r["rank"])
        style = "bold green" if r["priority"] == "now" else ""
        title = ("🎯 " if r["focus"] else "") + r["title"]
        table.add_row(str(i), f"[{style}]{rank}[/]" if style and rank else rank, _num(r["score"]),
                      _num(r["fresh"]), "🆕" if r["is_new"] else "", r["region"], r["age"],
                      title, r["company"], r["job_key"])
    console.print(table)
    earlier = len(rows) - len(shown)
    if earlier:
        console.print(f"[dim]{earlier} more posted earlier (still under a week) are in the report.[/dim]")
